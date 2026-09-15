#!/usr/bin/env python3
# /// script
# requires-python = ">=3.11"
# dependencies = ["python-pptx>=1.0", "lxml"]
# ///
"""Lint a generated .pptx against the confirmed-good Keynote/PowerPoint/
Google-Slides cross-compat rules and the course's slide-count cap.

This checks only the classes of breakage already confirmed and fixed
(see pptx_helpers.py's docstring) plus the deck-size constraint. It is
NOT a Keynote-import simulation -- the full deck's unresolved Keynote
failure (memory: project-session5-deck-keynote-fix) is out of scope.

Usage: lint_deck.py DECK.pptx [--min-slides 15] [--max-slides 20]
Exits 1 and prints one JSON object with an "issues" list (each entry
names the fix) if anything fails; exits 0 with {"ok": true} otherwise.
"""
import argparse
import json
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

ALLOWED_FONTS = {"Arial", "Helvetica Neue", "Times New Roman", "Courier New", None}


def check_fonts(prs) -> list[str]:
    # font.name can land on the paragraph's defRPr (paragraph.font.name --
    # what every pptx_helpers function sets) or on an individual run's rPr
    # (run.font.name). Checking only one leaves the other's bad fonts
    # invisible, so both are checked here.
    issues = []
    for si, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            text_frames = []
            if shape.has_text_frame:
                text_frames.append(shape.text_frame)
            elif getattr(shape, "has_table", False) and shape.has_table:
                text_frames.extend(cell.text_frame for row in shape.table.rows for cell in row.cells)
            for tf in text_frames:
                for para in tf.paragraphs:
                    if para.font.name not in ALLOWED_FONTS:
                        issues.append(
                            f"slide {si}: paragraph uses font '{para.font.name}', not in "
                            f"{sorted(f for f in ALLOWED_FONTS if f)} -- set font.name to 'Arial' "
                            f"(Windows-only faces make Keynote reject the file)"
                        )
                    for run in para.runs:
                        if run.font.name not in ALLOWED_FONTS:
                            issues.append(
                                f"slide {si}: run '{run.text[:30]!r}' uses font '{run.font.name}', not in "
                                f"{sorted(f for f in ALLOWED_FONTS if f)} -- set font.name to 'Arial' "
                                f"(Windows-only faces make Keynote reject the file)"
                            )
    return issues


def check_slide_count(prs, min_slides: int, max_slides: int) -> list[str]:
    n = len(prs.slides)
    if not (min_slides <= n <= max_slides):
        return [f"deck has {n} slides, outside the {min_slides}-{max_slides} cap -- trim or expand content"]
    return []


def check_sldsz_and_docprops(prs) -> list[str]:
    issues = []
    w, h = prs.slide_width, prs.slide_height
    expected_type = "screen16x9" if abs(w / h - 16 / 9) < abs(w / h - 4 / 3) else "screen4x3"
    actual_type = prs.part._element.find(qn("p:sldSz")).get("type")
    if actual_type != expected_type:
        issues.append(
            f"sldSz@type is '{actual_type}', expected '{expected_type}' for this aspect ratio -- "
            f"call pptx_helpers.finalize_and_save() instead of prs.save() directly"
        )

    app_part = next((p for p in prs.part.package.iter_parts() if p.partname == "/docProps/app.xml"), None)
    if app_part is not None:
        ns = {
            "ep": "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties",
            "vt": "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes",
        }
        app_xml = etree.fromstring(app_part.blob)
        slides_el = app_xml.find("ep:Slides", ns)
        recorded = slides_el.text if slides_el is not None else None
        if recorded != str(len(prs.slides)):
            issues.append(
                f"docProps/app.xml <Slides> says {recorded}, actual slide count is {len(prs.slides)} -- "
                f"call pptx_helpers.finalize_and_save() instead of prs.save() directly"
            )
    return issues


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("deck", help="path to the .pptx to lint")
    ap.add_argument("--min-slides", type=int, default=15)
    ap.add_argument("--max-slides", type=int, default=20)
    args = ap.parse_args()

    path = Path(args.deck)
    if not path.exists():
        print(json.dumps({"ok": False, "issues": [f"file not found: {path}"]}, indent=2))
        return 1

    prs = Presentation(str(path))
    issues = (
        check_slide_count(prs, args.min_slides, args.max_slides)
        + check_fonts(prs)
        + check_sldsz_and_docprops(prs)
    )

    print(json.dumps({"ok": not issues, "issues": issues}, indent=2))
    return 0 if not issues else 1


if __name__ == "__main__":
    sys.exit(main())
