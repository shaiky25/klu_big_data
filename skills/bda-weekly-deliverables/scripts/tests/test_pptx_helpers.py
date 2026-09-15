#!/usr/bin/env python3
"""Tests for pptx_helpers.py's Keynote-safety patching.
Run with: uv run --with pytest -m pytest test_pptx_helpers.py
(or plain `uv run test_pptx_helpers.py` for a lightweight self-check).
"""
import importlib.util
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn

SCRIPT = Path(__file__).resolve().parent.parent / "pptx_helpers.py"


def _load_module():
    spec = importlib.util.spec_from_file_location("pptx_helpers", SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def test_finalize_and_save_sets_16x9_sldsz():
    h = _load_module()
    prs = h.new_presentation()
    h.new_slide(prs)
    with tempfile.TemporaryDirectory() as d:
        out = str(Path(d) / "deck.pptx")
        h.finalize_and_save(prs, out)
        reopened = Presentation(out)
        assert reopened.part._element.find(qn("p:sldSz")).get("type") == "screen16x9"


def test_finalize_and_save_patches_docprops_slide_count():
    h = _load_module()
    prs = h.new_presentation()
    for _ in range(4):
        h.new_slide(prs)
    with tempfile.TemporaryDirectory() as d:
        out = str(Path(d) / "deck.pptx")
        h.finalize_and_save(prs, out)
        reopened = Presentation(out)
        app_part = next(p for p in reopened.part.package.iter_parts() if p.partname == "/docProps/app.xml")
        from lxml import etree
        ns = {"ep": "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"}
        app_xml = etree.fromstring(app_part.blob)
        assert app_xml.find("ep:Slides", ns).text == "4"


def test_add_card_uses_arial_only():
    h = _load_module()
    prs = h.new_presentation()
    slide = h.new_slide(prs)
    h.add_card(slide, h.Inches(1), h.Inches(1), h.Inches(3), h.Inches(2), "Header", ["a line"])
    shape = slide.shapes[-1]
    for para in shape.text_frame.paragraphs:
        assert para.font.name == "Arial"


if __name__ == "__main__":
    tests = [v for k, v in list(globals().items()) if k.startswith("test_")]
    failures = 0
    for t in tests:
        try:
            t()
            print(f"PASS {t.__name__}")
        except AssertionError as e:
            failures += 1
            print(f"FAIL {t.__name__}: {e}")
    sys.exit(1 if failures else 0)
