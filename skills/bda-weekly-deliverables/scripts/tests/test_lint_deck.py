#!/usr/bin/env python3
"""Tests for lint_deck.py's checks, including the paragraph-vs-run font
regression: pptx_helpers sets font.name on paragraph objects, which
python-pptx writes to the paragraph's defRPr, not any run's rPr -- a font
check that only reads run.font.name misses every real bad-font case.
Run with: uv run --with pytest -m pytest test_lint_deck.py
(or plain `uv run test_lint_deck.py` for a lightweight self-check).
"""
import importlib.util
import sys
import tempfile
from pathlib import Path

from pptx import Presentation

SCRIPTS_DIR = Path(__file__).resolve().parent.parent


def _load(name):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS_DIR / f"{name}.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def test_clean_deck_passes_all_checks():
    h = _load("pptx_helpers")
    lint = _load("lint_deck")
    prs = h.new_presentation()
    for i in range(16):
        slide = h.new_slide(prs)
        h.add_title(slide, f"Slide {i}")
        h.add_card(slide, h.Inches(0.6), h.Inches(1.6), h.Inches(5), h.Inches(2), "Header", ["a line"])
    with tempfile.TemporaryDirectory() as d:
        out = str(Path(d) / "deck.pptx")
        h.finalize_and_save(prs, out)
        reopened = Presentation(out)
        issues = (
            lint.check_slide_count(reopened, 15, 20)
            + lint.check_fonts(reopened)
            + lint.check_sldsz_and_docprops(reopened)
        )
        assert issues == []


def test_bad_paragraph_font_is_caught():
    h = _load("pptx_helpers")
    lint = _load("lint_deck")
    prs = h.new_presentation()
    slide = h.new_slide(prs)
    box = slide.shapes.add_textbox(h.Inches(1), h.Inches(1), h.Inches(3), h.Inches(1))
    p = box.text_frame.paragraphs[0]
    p.text = "bad"
    p.font.name = "Segoe UI"  # sets defRPr, not run.font.name -- the exact regression case
    issues = lint.check_fonts(prs)
    assert any("Segoe UI" in i for i in issues)


def test_slide_count_outside_cap_is_caught():
    h = _load("pptx_helpers")
    lint = _load("lint_deck")
    prs = h.new_presentation()
    h.new_slide(prs)
    issues = lint.check_slide_count(prs, 15, 20)
    assert len(issues) == 1


def test_unpatched_save_is_caught():
    h = _load("pptx_helpers")
    lint = _load("lint_deck")
    prs = h.new_presentation()
    h.new_slide(prs)
    with tempfile.TemporaryDirectory() as d:
        out = str(Path(d) / "deck.pptx")
        prs.save(out)  # raw save, skips finalize_and_save's patches
        reopened = Presentation(out)
        issues = lint.check_sldsz_and_docprops(reopened)
        assert len(issues) == 2  # sldSz@type wrong, docProps slide count stale


if __name__ == "__main__":
    tests = [v for k, v in list(globals().items()) if k.startswith("test_")]
    failures = 0
    for t in tests:
        try:
            t()
            print(f"PASS {t.__name__}")
        except AssertionError as e:
            failures += 1
            print(f"FAIL {t.__name__}: {e}")
    sys.exit(1 if failures else 0)
